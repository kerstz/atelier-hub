import csv
import os
import re
import shutil
from pathlib import Path
from typing import Callable

from . import translate
from .batch import Translator, numbered
from .config import cfg, data_dir, output_dir
from .jobs import Ctx, manager, register
from .proc import run_cmd
from .srt import flat

PCT_RE = re.compile(r"(\d{1,3})(?:\.\d+)?\s*%")
PAGES_RE = re.compile(r"^[0-9,\- ]{1,60}$")
CONVERT = {".doc": "docx", ".odt": "docx", ".rtf": "docx", ".ppt": "pptx", ".odp": "pptx", ".xls": "xlsx", ".ods": "xlsx"}
Slots = list[tuple[str, Callable[[str], None]]]


def load_docx(path: Path) -> tuple[Slots, Callable[[Path], None]]:
    import docx
    from docx.oxml.ns import qn

    doc = docx.Document(str(path))
    roots = [doc.element.body]
    for rel in doc.part.rels.values():
        if rel.reltype.endswith(("/header", "/footer", "/footnotes", "/endnotes")):
            roots.append(rel.target_part.element)
    w_p, w_t = qn("w:p"), qn("w:t")

    def owner(t):
        e = t.getparent()
        while e is not None and e.tag != w_p:
            e = e.getparent()
        return e

    def setter(ts):
        def set_(text: str) -> None:
            ts[0].text = text
            ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            for t in ts[1:]:
                t.text = ""
        return set_

    slots: Slots = []
    for root in roots:
        for p in root.iter(w_p):
            ts = [t for t in p.iter(w_t) if owner(t) is p]
            text = "".join(t.text or "" for t in ts)
            if ts and re.search(r"\w", text):
                slots.append((text, setter(ts)))
    return slots, lambda dest: doc.save(str(dest))


def load_pptx(path: Path) -> tuple[Slots, Callable[[Path], None]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))

    def shapes(coll):
        for sh in coll:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from shapes(sh.shapes)
            else:
                yield sh

    frames = []
    for slide in prs.slides:
        for sh in shapes(slide.shapes):
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                frames.append(sh.text_frame)
            if getattr(sh, "has_table", False) and sh.has_table:
                frames += [c.text_frame for row in sh.table.rows for c in row.cells]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            frames.append(slide.notes_slide.notes_text_frame)

    def setter(runs):
        def set_(text: str) -> None:
            runs[0].text = text
            for r in runs[1:]:
                r.text = ""
        return set_

    slots: Slots = []
    for tf in frames:
        for para in tf.paragraphs:
            runs = list(para.runs)
            text = "".join(r.text for r in runs)
            if runs and re.search(r"\w", text):
                slots.append((text, setter(runs)))
    return slots, lambda dest: prs.save(str(dest))


def load_xlsx(path: Path) -> tuple[Slots, Callable[[Path], None]]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path))

    def setter(cell):
        def set_(text: str) -> None:
            cell.value = text
        return set_

    slots: Slots = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and not v.startswith("=") and re.search(r"\w", v):
                    slots.append((v, setter(cell)))
    return slots, lambda dest: wb.save(str(dest))


LOADERS = {".docx": load_docx, ".pptx": load_pptx, ".xlsx": load_xlsx}


async def translate_texts(ctx: Ctx, texts: list[str], src, tgt, glossary, lo: float, hi: float) -> dict[str, str]:
    uniq = list(dict.fromkeys(texts))
    items = list(enumerate(uniq, 1))
    limit = int(cfg["translate"]["chunk_chars"])
    batches: list[list[tuple[int, str]]] = []
    cur: list[tuple[int, str]] = []
    size = 0
    for it in items:
        if cur and (size + len(it[1]) > limit or len(cur) >= 40):
            batches.append(cur)
            cur, size = [], 0
        cur.append(it)
        size += len(it[1])
    if cur:
        batches.append(cur)
    tr, res, prev = Translator(src, tgt, glossary), {}, []
    ctx.set(lo, "attente GPU…")
    async with manager.gpu_lock:
        for bi, b in enumerate(batches):
            got = await tr.batch(b, prev)
            res.update(got)
            prev = [
                {"role": "user", "content": tr.prompt(numbered(b))},
                {"role": "assistant", "content": numbered([(n, got[n]) for n, _ in b])},
            ]
            ctx.set(lo + (hi - lo) * (bi + 1) / len(batches), f"traduction {bi + 1}/{len(batches)}")
    return {t: res[n] for n, t in items}


async def do_office(ctx: Ctx, path: Path, ext: str, name: str, src, tgt, glossary, out: Path) -> Path:
    ctx.set(0.02, "lecture…")
    slots, save = LOADERS[ext](path)
    texts = [flat(t) for t, _ in slots]
    if not texts:
        raise RuntimeError("aucun texte à traduire")
    done = await translate_texts(ctx, texts, src, tgt, glossary, 0.05, 0.95)
    for (_, setter), t in zip(slots, texts):
        if t in done:
            setter(done[t])
    dest = out / f"{name}.{tgt}{ext}"
    save(dest)
    return dest


async def do_pdf(ctx: Ctx, path: Path, name: str, src, tgt, glossary, pages: str, work: Path, out: Path) -> list[Path]:
    if src == "auto":
        raise RuntimeError("langue source requise pour un PDF")
    d, ol = cfg["docs"], cfg["ollama"]
    res_dir = work / "out"
    res_dir.mkdir()
    pdf = work / f"{name}.pdf"
    shutil.copy(path, pdf)
    argv = [
        "pdf2zh_next", "--ollama", "--ollama-model", d["pdf_model"] or ol["translate_model"],
        "--ollama-host", ol["url"], "--lang-in", src, "--lang-out", tgt, "--output", str(res_dir),
        "--qps", str(d["qps"]), "--watermark-output-mode", "no_watermark", "--disable-config-auto-save",
    ]
    if not d["auto_glossary"]:
        argv.append("--no-auto-extract-glossary")
    if pages:
        if not PAGES_RE.match(pages):
            raise RuntimeError("pages invalides")
        argv += ["--pages", pages.replace(" ", "")]
    if glossary:
        gfile = work / "glossary.csv"
        with open(gfile, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["source", "target"])
            w.writerows(glossary)
        argv += ["--glossaries", str(gfile)]
    argv.append(str(pdf))

    def feed(line: str) -> None:
        if m := PCT_RE.search(line):
            ctx.set(min(int(m[1]), 99) / 100, line[:70])
        else:
            ctx.set(message=line[:70])

    ctx.set(0, "attente GPU…")
    async with manager.gpu_lock:
        ctx.set(0, "traduction PDF…")
        env = {**os.environ, "PYTHONUNBUFFERED": "1", "NO_COLOR": "1", "TERM": "dumb"}
        await run_cmd(argv, feed, env=env)
    pdfs = sorted(res_dir.glob("*.pdf"))
    if not pdfs:
        raise RuntimeError("pdf2zh: aucune sortie")
    dests = []
    for f in pdfs:
        dests.append(out / f.name.replace(".no_watermark", ""))
        shutil.move(f, dests[-1])
    return dests


async def run(ctx: Ctx, p: dict) -> str:
    path = Path(p.get("path") or "")
    src, tgt = p.get("src", "en"), p.get("tgt", "fr")
    if not path.is_file():
        raise RuntimeError("fichier introuvable")
    if tgt not in translate.LANGS or (src != "auto" and src not in translate.LANGS):
        raise RuntimeError("langue inconnue")
    ext = path.suffix.lower()
    if ext != ".pdf" and ext not in LOADERS and ext not in CONVERT:
        raise RuntimeError(f"format non supporté: {ext}")
    glossary = translate.parse_glossary(p.get("glossary") or "")
    uploads = (data_dir() / "uploads").resolve()
    name = re.sub(r"^[0-9a-f]{8}-", "", path.name) if path.resolve().is_relative_to(uploads) else path.name
    name = re.sub(r'[\\/:*?"<>|]+', "_", Path(name).stem).strip(" .")[:120] or "document"
    work = data_dir() / "tmp" / str(ctx.id)
    shutil.rmtree(work, ignore_errors=True)
    work.mkdir(parents=True)
    out = output_dir("docs")
    out.mkdir(parents=True, exist_ok=True)
    try:
        if ext == ".pdf":
            dests = await do_pdf(ctx, path, name, src, tgt, glossary, (p.get("pages") or "").strip(), work, out)
            ctx.set(1.0, ", ".join(d.name for d in dests))
            return str(out)
        if ext in CONVERT:
            ctx.set(0, "conversion LibreOffice…")
            await run_cmd(
                ["soffice", "--headless", "--convert-to", CONVERT[ext], "--outdir", str(work), str(path)]
            )
            conv = work / f"{path.stem}.{CONVERT[ext]}"
            if not conv.is_file():
                raise RuntimeError("conversion échouée")
            path, ext = conv, "." + CONVERT[ext]
        dest = await do_office(ctx, path, ext, name, src, tgt, glossary, out)
        ctx.set(1.0, dest.name)
        return str(out)
    finally:
        shutil.rmtree(work, ignore_errors=True)
        orig = Path(p["path"])
        if orig.resolve().is_relative_to(uploads):
            orig.unlink(missing_ok=True)


register("document", run)
