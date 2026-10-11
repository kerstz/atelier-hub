import copy
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from .inline import Fmt, Seg, encode

XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


@dataclass
class Slot:
    text: str
    put: Callable[[str], None]
    tagged: str | None = None
    links: frozenset[int] = field(default_factory=frozenset)
    put_segs: Callable[[list[Seg]], None] | None = None


def _pick(templates: list, f: Fmt):
    """Template run closest to wanted format: exact, then same link, then first."""
    for r, tf in templates:
        if tf == f:
            return r
    for r, tf in templates:
        if tf.link == f.link:
            return r
    return templates[0][0]


def load_docx(path: Path) -> tuple[list[Slot], Callable[[Path], None]]:
    import docx
    from docx.oxml.ns import qn
    from docx.text.run import Run

    doc = docx.Document(str(path))
    roots = [doc.element.body]
    for rel in doc.part.rels.values():
        if rel.reltype.endswith(("/header", "/footer", "/footnotes", "/endnotes")):
            roots.append(rel.target_part.element)
    w_p, w_r, w_t, w_hl, w_rpr = (qn(f"w:{n}") for n in ("p", "r", "t", "hyperlink", "rPr"))
    w_fld = (qn("w:fldChar"), qn("w:fldSimple"))

    def owner(t):
        e = t.getparent()
        while e is not None and e.tag != w_p:
            e = e.getparent()
        return e

    def flag(r, tag: str) -> bool:
        rpr = r.find(w_rpr)
        el = rpr.find(qn(tag)) if rpr is not None else None
        return el is not None and el.get(qn("w:val")) not in ("0", "false", "off")

    def setter(ts):
        def set_(text: str) -> None:
            ts[0].text = text
            ts[0].set(XML_SPACE, "preserve")
            for t in ts[1:]:
                t.text = ""
        return set_

    def rich(p, ts):
        """Segments + rebuild fn when every text run is a plain run or sits in a direct hyperlink."""
        if any(True for _ in p.iter(*w_fld)):
            return None
        links: dict = {}
        tops, templates, segs = [], [], []
        for r in dict.fromkeys(t.getparent() for t in ts):
            par = r.getparent()
            if r.tag != w_r or any(c.tag not in (w_rpr, w_t) for c in r):
                return None
            if par is p:
                top, link = r, 0
            elif par.tag == w_hl and par.getparent() is p:
                top, link = par, links.setdefault(par, len(links) + 1)
            else:
                return None
            if top not in tops:
                tops.append(top)
            fmt = Fmt(flag(r, "w:b"), flag(r, "w:i"), link)
            templates.append((r, fmt))
            segs.append(("".join(t.text or "" for t in r.findall(w_t)), fmt))
        by_id = {i: el for el, i in links.items()}

        def put(out: list[Seg]) -> None:
            new = []
            for text, f in out:
                run = copy.deepcopy(_pick(templates, f))
                for t in run.findall(w_t):
                    run.remove(t)
                t = run.makeelement(w_t, {})
                t.text = text
                t.set(XML_SPACE, "preserve")
                run.append(t)
                rr = Run(run, None)
                if bool(rr.bold) != f.b:
                    rr.bold = f.b
                if bool(rr.italic) != f.i:
                    rr.italic = f.i
                if f.link:
                    hl = copy.deepcopy(by_id[f.link])
                    for c in list(hl):
                        hl.remove(c)
                    hl.append(run)
                    run = hl
                new.append(run)
            idx = p.index(tops[0])
            for top in tops:
                p.remove(top)
            for k, el in enumerate(new):
                p.insert(idx + k, el)

        return segs, frozenset(links.values()), put

    slots: list[Slot] = []
    for root in roots:
        for p in root.iter(w_p):
            ts = [t for t in p.iter(w_t) if owner(t) is p]
            text = "".join(t.text or "" for t in ts)
            if not (ts and re.search(r"\w", text)):
                continue
            slot = Slot(text, setter(ts))
            if r := rich(p, ts):
                slot.tagged = encode(r[0])
                slot.links, slot.put_segs = r[1], r[2]
            slots.append(slot)
    return slots, lambda dest: doc.save(str(dest))


def load_pptx(path: Path) -> tuple[list[Slot], Callable[[Path], None]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.oxml.ns import qn
    from pptx.text.text import _Run

    prs = Presentation(str(path))
    a_br, a_fld, a_rpr, a_hl = (qn(f"a:{n}") for n in ("br", "fld", "rPr", "hlinkClick"))

    def shapes(coll):
        for sh in coll:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from shapes(sh.shapes)
            else:
                yield sh

    frames = []
    for slide in prs.slides:
        for sh in shapes(slide.shapes):
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                frames.append(sh.text_frame)
            if getattr(sh, "has_table", False) and sh.has_table:
                frames += [c.text_frame for row in sh.table.rows for c in row.cells]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            frames.append(slide.notes_slide.notes_text_frame)

    def setter(runs):
        def set_(text: str) -> None:
            runs[0].text = text
            for r in runs[1:]:
                r.text = ""
        return set_

    def rich(para, runs):
        p = para._p
        if p.find(a_br) is not None or p.find(a_fld) is not None:
            return None
        links: dict = {}
        templates, segs = [], []
        for r in runs:
            rpr = r._r.find(a_rpr)
            hl = rpr.find(a_hl) if rpr is not None else None
            key = (hl.get(qn("r:id")), hl.get("action")) if hl is not None else None
            link = links.setdefault(key, len(links) + 1) if key else 0
            fmt = Fmt(bool(r.font.bold), bool(r.font.italic), link)
            templates.append((r._r, fmt))
            segs.append((r.text, fmt))

        def put(out: list[Seg]) -> None:
            first = runs[0]._r
            for text, f in out:
                el = copy.deepcopy(_pick(templates, f))
                rr = _Run(el, para)
                rr.text = text
                rpr = el.find(a_rpr)
                if not f.link and rpr is not None:
                    for h in rpr.findall(a_hl):
                        rpr.remove(h)
                if bool(rr.font.bold) != f.b:
                    rr.font.bold = f.b
                if bool(rr.font.italic) != f.i:
                    rr.font.italic = f.i
                first.addprevious(el)
            for r in runs:
                p.remove(r._r)

        return segs, frozenset(links.values()), put

    slots: list[Slot] = []
    for tf in frames:
        for para in tf.paragraphs:
            runs = list(para.runs)
            text = "".join(r.text for r in runs)
            if not (runs and re.search(r"\w", text)):
                continue
            slot = Slot(text, setter(runs))
            if r := rich(para, runs):
                slot.tagged = encode(r[0])
                slot.links, slot.put_segs = r[1], r[2]
            slots.append(slot)
    return slots, lambda dest: prs.save(str(dest))


def load_xlsx(path: Path) -> tuple[list[Slot], Callable[[Path], None]]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path))

    def setter(cell):
        def set_(text: str) -> None:
            cell.value = text
        return set_

    slots: list[Slot] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and not v.startswith("=") and re.search(r"\w", v):
                    slots.append(Slot(v, setter(cell)))
    return slots, lambda dest: wb.save(str(dest))


LOADERS = {".docx": load_docx, ".pptx": load_pptx, ".xlsx": load_xlsx}
